"""
report.py — 계산 결과를 의뢰인이 받아 볼 보고서 엑셀로 저장
합계 칸은 엑셀 수식(SUM)으로 넣어서, 의뢰인이 표를 고쳐도 합계가 따라 바뀌게 함.
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
HEAD = PatternFill("solid", fgColor="EDEFF2")
WON = '#,##0;[Red]-#,##0'


def _table(ws, columns, rows, start_row=1, money=(), freeze=True):
    """표 하나 쓰기: 열 이름 줄 + 데이터. money 열은 천 단위 쉼표, 음수는 빨강."""
    for j, (name, width) in enumerate(columns, 1):
        c = ws.cell(start_row, j, name); c.font = Font(name=FONT, bold=True); c.fill = HEAD
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, row in enumerate(rows, start_row + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.font = Font(name=FONT)
            if columns[j - 1][0] in money: c.number_format = WON
    if freeze:   # 열 이름 줄 고정. 요약 시트처럼 표가 중간에 있으면 고정하지 않음
        ws.freeze_panes = ws.cell(start_row + 1, 1)
    return start_row + len(rows)   # 마지막 데이터 줄 번호


def save_report(result, path):
    wb = Workbook()
    t = result["totals"]

    # 1) 요약
    s = wb.active; s.title = "요약"
    s["A1"] = "무숙녀 스마트스토어·쿠팡 정산 보고서"; s["A1"].font = Font(name=FONT, bold=True, size=15)
    s["A2"] = f"주문 기간 {result['period'][0]} ~ {result['period'][1]}"; s["A2"].font = Font(name=FONT, color="666666")
    kpi = [("판매 수량", t["판매수량"], "0"), ("정산금액", t["정산금액"], WON), ("순이익", t["순이익"], WON),
           ("순이익률", t["순이익"] / t["정산금액"] if t["정산금액"] else 0, "0.0%"),
           ("받을 돈 (정산 누락)", t["받을 돈"], WON), ("금액 오류", t["오류 건수"], "0\"건\"")]
    for i, (k, v, fmt) in enumerate(kpi, 4):
        s.cell(i, 1, k).font = Font(name=FONT, color="666666")
        c = s.cell(i, 2, v); c.font = Font(name=FONT, bold=True); c.number_format = fmt
    s.cell(11, 1, "쇼핑몰별").font = Font(name=FONT, bold=True)
    last = _table(s, [("쇼핑몰", 16), ("정산금액", 16), ("순이익", 16)],
                  [(mall, v["정산금액"], v["순이익"]) for mall, v in t["쇼핑몰별"].items()], start_row=12, money=("정산금액", "순이익"), freeze=False)
    s.column_dimensions["A"].width = 22
    if result["warnings"]:
        s.cell(last + 2, 1, "확인할 것").font = Font(name=FONT, bold=True, color="B42335")
        for i, w in enumerate(result["warnings"], last + 3):
            c = s.cell(i, 1, f"• {w}"); c.font = Font(name=FONT); s.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)

    # 2) 상품별 순이익 (+ 합계 수식, 막대 그래프)
    p = wb.create_sheet("상품별 순이익")
    cols = [("쇼핑몰", 12), ("상품코드", 9), ("상품명", 24), ("판매수량", 9), ("정산금액", 13), ("원가합계", 12),
            ("포장비", 10), ("배송비", 10), ("광고비", 11), ("반품비용", 11), ("순이익", 13)]
    data = result["summary"][[c for c, _ in cols]].values.tolist()
    last = _table(p, cols, data, money=("정산금액", "원가합계", "포장비", "배송비", "광고비", "반품비용", "순이익"))
    total_row = last + 1
    p.cell(total_row, 1, "합계").font = Font(name=FONT, bold=True)
    for j in range(4, len(cols) + 1):
        col = get_column_letter(j)
        c = p.cell(total_row, j, f"=SUM({col}2:{col}{last})"); c.font = Font(name=FONT, bold=True); c.number_format = WON
    chart = BarChart(); chart.type = "bar"; chart.title = "쇼핑몰·상품별 순이익"; chart.height = 12; chart.width = 18
    chart.add_data(Reference(p, min_col=11, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(p, min_col=3, min_row=2, max_row=last)); chart.legend = None
    p.add_chart(chart, "M2")

    # 3) 받을 돈과 오류
    e = wb.create_sheet("받을 돈과 오류")
    rows = result["issues"].sort_values(["종류", "쇼핑몰"]).values.tolist()
    _table(e, [("종류", 12), ("쇼핑몰", 12), ("상품주문번호", 14), ("상품", 34), ("내용", 60), ("금액", 12)], rows, money=("금액",))

    # 4) 줄별 계산 (어떻게 계산됐는지 한 줄씩 확인)
    L = result["lines"]
    lc = [("쇼핑몰", 12), ("주문번호", 14), ("상품주문번호", 14), ("상품코드", 9), ("상태", 9), ("배송비구분", 10), ("수량", 6),
          ("정산금액", 11), ("원가합계", 11), ("포장비", 9), ("배송비", 9), ("광고비", 9), ("반품비용", 10), ("순이익", 11), ("정산누락", 9)]
    rows = [[("누락" if v is True else "" if v is False else v) for v in r] for r in L[[c for c, _ in lc]].values.tolist()]
    _table(wb.create_sheet("줄별 계산"), lc, rows, money=("정산금액", "원가합계", "포장비", "배송비", "광고비", "반품비용", "순이익"))

    # 5) 계산 규칙 (보고서만 받아 본 사람도 알 수 있게)
    r = wb.create_sheet("계산 규칙")
    rules = [("순이익", "정산금액 - 원가 x 수량 - 포장비 - 판매자 부담 배송비 몫 - 광고비 몫 - 반품 비용"),
             ("반품", "정산 0, 원가 0 (상품이 돌아옴). 포장비, 배송비, 반품 비용은 손실"), ("취소", "계산에서 뺌"),
             ("정산 누락", "순이익에서 빼고 '받을 돈'으로 따로 보여 줌"), ("광고비", "쇼핑몰·상품별 광고비를 팔린 수량대로 나눔"),
             ("수수료 확인", "그 달 쇼핑몰별 수수료율 가운뎃값보다 1%p 넘게 다르면 오류로 표시"),
             ("금액", "모두 원 단위 정수. 나눌 때 남는 원은 첫 줄에")]
    _table(r, [("항목", 14), ("규칙", 90)], rules)
    wb.save(path)
